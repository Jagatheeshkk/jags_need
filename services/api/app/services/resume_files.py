"""Validation and text extraction for uploaded resumes (PDF and DOCX only).

Uploaded bytes are untrusted. Nothing is written to disk: files are checked in memory,
then stored in PostgreSQL. The declared type must match the sniffed content, DOCX archives are
checked for zip bombs before parsing, and encrypted PDFs are refused.
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
ALLOWED = {".pdf": PDF_MIME, ".docx": DOCX_MIME}

MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED = 30 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_TEXT = 100_000


class UploadRejected(ValueError):
    pass


@dataclass
class CheckedFile:
    file_name: str
    extension: str
    mime_type: str
    data: bytes
    text: str


def safe_filename(name: str | None, extension: str) -> str:
    """Strip any path (both separators), control and unusual characters; force the real extension."""
    base = re.split(r"[\\/]", name or "")[-1]
    base = unicodedata.normalize("NFKC", base)
    stem = base.rsplit(".", 1)[0] if "." in base else base
    stem = re.sub(r"[^\w\- ]+", "_", stem).strip(" ._")[:100] or "resume"
    return f"{stem}{extension}"


def sniff(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return ".pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile as exc:
            raise UploadRejected("The file is not a valid DOCX document.") from exc
        if "word/document.xml" in names and "[Content_Types].xml" in names:
            return ".docx"
    raise UploadRejected("Only PDF and Word (.docx) resumes are accepted.")


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = z.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise UploadRejected("The DOCX file has too many parts.")
        total = sum(i.file_size for i in infos)
        if total > MAX_ZIP_UNCOMPRESSED:
            raise UploadRejected("The DOCX file expands to an unreasonable size and was refused.")


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT]


def extract_text(extension: str, data: bytes) -> str:
    try:
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise UploadRejected("Password-protected PDFs cannot be read. Upload an unprotected copy.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise UploadRejected(f"The PDF has more than {MAX_PDF_PAGES} pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            import docx

            _check_zip(data)
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
    except UploadRejected:
        raise
    except Exception as exc:  # noqa: BLE001 - any parser failure means an unreadable file
        raise UploadRejected("The file could not be read. It may be damaged.") from exc
    return _clean(text)


def check_upload(file_name: str | None, data: bytes, max_bytes: int) -> CheckedFile:
    if not data:
        raise UploadRejected("The file is empty.")
    if len(data) > max_bytes:
        raise UploadRejected(f"The file is larger than {max_bytes // 1024} KB.")
    declared = ("." + file_name.rsplit(".", 1)[-1].lower()) if file_name and "." in file_name else ""
    if declared not in ALLOWED:
        raise UploadRejected("Only PDF and Word (.docx) resumes are accepted.")
    actual = sniff(data)
    if actual != declared:
        raise UploadRejected(f"The file content is {actual.upper()[1:]}, but its name ends in {declared}.")
    text = extract_text(actual, data)
    return CheckedFile(safe_filename(file_name, actual), actual, ALLOWED[actual], data, text)
