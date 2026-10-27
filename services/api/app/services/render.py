"""Render generated content to DOCX (python-docx) and PDF (ReportLab), locally and for free.

All text is inserted as plain text. For PDF, ReportLab Paragraphs interpret markup, so every
value is XML-escaped first.
"""

import io
from typing import Any, Literal
from xml.sax.saxutils import escape

Format = Literal["docx", "pdf"]


def _contact_line(c: dict[str, Any]) -> str:
    parts = [c.get("email"), c.get("phone"), c.get("location")] + list((c.get("links") or {}).values())
    return "  |  ".join(p for p in parts if p)


def _resume_blocks(content: dict[str, Any]) -> list[tuple[str, str]]:
    """Flatten a resume into (style, text) blocks shared by both renderers."""
    b: list[tuple[str, str]] = [("title", content.get("name") or "Resume")]
    if content.get("headline"):
        b.append(("subtitle", content["headline"]))
    contact = _contact_line(content.get("contact") or {})
    if contact:
        b.append(("small", contact))
    if content.get("summary"):
        b.append(("h2", "Professional Summary"))
        b += [("p", para) for para in content["summary"].split("\n\n") if para.strip()]
    if content.get("skills"):
        b.append(("h2", "Skills"))
        b.append(("p", ", ".join(content["skills"])))
    if content.get("experience"):
        b.append(("h2", "Experience"))
        for e in content["experience"]:
            where = f" — {e['location']}" if e.get("location") else ""
            b.append(("h3", f"{e['title']}, {e['company']}{where}"))
            b.append(("small", e["dates"]))
            if e.get("description"):
                b.append(("p", e["description"]))
            b += [("bullet", a) for a in e.get("achievements") or []]
    if content.get("education"):
        b.append(("h2", "Education"))
        for ed in content["education"]:
            years = "–".join(str(y) for y in (ed.get("start_year"), ed.get("end_year")) if y)
            field = f", {ed['field_of_study']}" if ed.get("field_of_study") else ""
            b.append(("p", f"{ed['degree']}{field}, {ed['institution']}" + (f" ({years})" if years else "")))
    if content.get("certifications"):
        b.append(("h2", "Certifications"))
        for c in content["certifications"]:
            extra = ", ".join(str(x) for x in (c.get("issuer"), c.get("year")) if x)
            b.append(("bullet", c["name"] + (f" ({extra})" if extra else "")))
    return b


def _letter_blocks(content: dict[str, Any]) -> list[tuple[str, str]]:
    b: list[tuple[str, str]] = []
    contact = _contact_line(content.get("contact") or {})
    if content.get("name"):
        b.append(("h3", content["name"]))
    if contact:
        b.append(("small", contact))
    b.append(("p", content["greeting"]))
    b += [("p", p) for p in content["paragraphs"]]
    b.append(("p", content["sign_off"]))
    if content.get("name"):
        b.append(("p", content["name"]))
    return b


def _answers_blocks(content: dict[str, Any]) -> list[tuple[str, str]]:
    target = content.get("target")
    b: list[tuple[str, str]] = [("title", "Screening answer drafts")]
    if target:
        b.append(("subtitle", f"{target['title']} — {target['company']}"))
    b.append(("small", "Drafts only. Items marked REVIEW need your confirmation before use."))
    for a in content["answers"]:
        b.append(("h3", a["question"]))
        flag = "REVIEW: " if a["needs_confirmation"] else ""
        b.append(("p", f"{flag}{a['answer'] or '(no answer: ' + a['reason'] + ')'}"))
        if a["needs_confirmation"] and a["answer"]:
            b.append(("small", a["reason"]))
    return b


def _docx(blocks: list[tuple[str, str]]) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.core_properties.author = ""
    d.core_properties.comments = "Generated locally by AI Job Agent"
    for style, text in blocks:
        if style == "title":
            d.add_heading(text, level=0)
        elif style == "h2":
            d.add_heading(text, level=1)
        elif style == "h3":
            d.add_heading(text, level=2)
        elif style == "bullet":
            d.add_paragraph(text, style="List Bullet")
        else:
            p = d.add_paragraph(text)
            if style in {"small", "subtitle"}:
                for run in p.runs:
                    run.font.size = Pt(9 if style == "small" else 12)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(blocks: list[tuple[str, str]], title: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    style_map = {
        "title": styles["Title"],
        "subtitle": styles["Heading3"],
        "h2": styles["Heading2"],
        "h3": styles["Heading4"],
        "small": styles["Italic"],
        "p": styles["BodyText"],
    }
    story: list[Any] = []
    bullets: list[Any] = []

    def flush() -> None:
        if bullets:
            story.append(ListFlowable(list(bullets), bulletType="bullet", leftIndent=12))
            bullets.clear()

    for style, text in blocks:
        safe = escape(text).replace("\n", "<br/>")
        if style == "bullet":
            bullets.append(ListItem(Paragraph(safe, styles["BodyText"])))
            continue
        flush()
        story.append(Paragraph(safe, style_map.get(style, styles["BodyText"])))
        if style in {"p", "small"}:
            story.append(Spacer(1, 2 * mm))
    flush()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        title=title,
        author="",
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    doc.build(story)
    return buf.getvalue()


def render_resume(content: dict[str, Any], fmt: Format) -> bytes:
    blocks = _resume_blocks(content)
    return _docx(blocks) if fmt == "docx" else _pdf(blocks, content.get("name") or "Resume")


def render_document(content: dict[str, Any], fmt: Format) -> bytes:
    blocks = _letter_blocks(content) if content["type"] == "cover_letter" else _answers_blocks(content)
    title = "Cover letter" if content["type"] == "cover_letter" else "Screening answers"
    return _docx(blocks) if fmt == "docx" else _pdf(blocks, title)
